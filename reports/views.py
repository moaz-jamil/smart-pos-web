import io
from decimal import Decimal
from django.shortcuts import render
from django.http import HttpResponse
from django.contrib.auth.decorators import login_required
from django.utils import timezone
from django.db.models import Sum

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from sales.models import Sale
from products.models import Product
from expenses.models import Expense
from purchases.models import Purchase

@login_required
def reports_overview(request):
    business = request.user.business
    sales_total = Sale.objects.filter(business=business).aggregate(total=Sum('total_amount'))['total'] or Decimal('0.00')
    purchases_total = Purchase.objects.filter(business=business).aggregate(total=Sum('total_amount'))['total'] or Decimal('0.00')
    expenses_total = Expense.objects.filter(business=business).aggregate(total=Sum('amount'))['total'] or Decimal('0.00')
    net_profit = sales_total - purchases_total - expenses_total

    context = {
        'sales_total': sales_total,
        'purchases_total': purchases_total,
        'expenses_total': expenses_total,
        'net_profit': net_profit,
    }
    return render(request, 'reports/reports_overview.html', context)


@login_required
def export_excel_sales(request):
    business = request.user.business
    sales = Sale.objects.filter(business=business).order_by('-created_at') if business else Sale.objects.all().order_by('-created_at')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales Report"

    # Formatting Header
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    headers = ["Invoice #", "Date", "Customer", "Cashier", "Payment Method", "Total Amount ($)", "Status"]
    ws.append(headers)

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for sale in sales:
        ws.append([
            sale.invoice_number,
            sale.created_at.strftime("%Y-%m-%d %H:%M"),
            sale.customer.name if sale.customer else "Walk-in",
            sale.cashier.username if sale.cashier else "System",
            sale.get_payment_method_display(),
            float(sale.total_amount),
            sale.payment_status
        ])

    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response['Content-Disposition'] = f'attachment; filename="smartpos_sales_report_{timezone.now().strftime("%Y%m%d")}.xlsx"'
    wb.save(response)
    return response


@login_required
def export_pdf_sales(request):
    business = request.user.business
    sales = Sale.objects.filter(business=business).order_by('-created_at')[:50] if business else Sale.objects.all().order_by('-created_at')[:50]

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
    story = []

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(name='TitleStyle', parent=styles['Heading1'], fontSize=18, textColor=colors.HexColor('#4F46E5'))

    story.append(Paragraph(f"Sales Report - {business.name if business else 'SmartPOS Pro'}", title_style))
    story.append(Paragraph(f"Generated Date: {timezone.now().strftime('%Y-%m-%d %H:%M')}", styles['Normal']))
    story.append(Spacer(1, 15))

    data = [["Invoice #", "Date", "Customer", "Method", "Total ($)", "Status"]]
    for s in sales:
        data.append([
            s.invoice_number,
            s.created_at.strftime("%Y-%m-%d"),
            s.customer.name if s.customer else "Walk-in",
            s.get_payment_method_display(),
            f"${s.total_amount:.2f}",
            s.payment_status
        ])

    table = Table(data, colWidths=[110, 80, 130, 80, 80, 70])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E293B')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('ALIGN', (4, 1), (4, -1), 'RIGHT'),
    ]))

    story.append(table)
    doc.build(story)

    buffer.seek(0)
    response = HttpResponse(buffer, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="smartpos_sales_report_{timezone.now().strftime("%Y%m%d")}.pdf"'
    return response


@login_required
def export_excel_inventory(request):
    business = request.user.business
    products = Product.objects.filter(business=business) if business else Product.objects.all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventory Audit"

    header_fill = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    headers = ["Barcode", "SKU", "Product Name", "Category", "Cost Price ($)", "Selling Price ($)", "Stock Qty", "Low Stock Alert"]
    ws.append(headers)

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font

    for p in products:
        ws.append([
            p.barcode,
            p.sku,
            p.name,
            p.category.name if p.category else "-",
            float(p.purchase_price),
            float(p.selling_price),
            p.stock_quantity,
            "YES" if p.is_low_stock else "NO"
        ])

    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response['Content-Disposition'] = f'attachment; filename="smartpos_inventory_{timezone.now().strftime("%Y%m%d")}.xlsx"'
    wb.save(response)
    return response
